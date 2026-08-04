from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

source = Path(r"C:\Users\Владимир\Desktop\Сайты\Передача\Презентация Peredacha\Раб.файлы для видео\Сценарий_видеопрезентации_CRM_Передача.docx")
doc = Document(source)

def value(v):
    return None if v is None else str(v)

for idx in [2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 13, 14, 53, 54, 55, 56, 61, 62]:
    p = doc.paragraphs[idx - 1]
    pf = p.paragraph_format
    print(f"P{idx:03d} text={p.text[:70]!r} style={p.style.name}")
    print("  paragraph", {
        "align": value(p.alignment),
        "before": value(pf.space_before),
        "after": value(pf.space_after),
        "line": value(pf.line_spacing),
        "left": value(pf.left_indent),
        "right": value(pf.right_indent),
        "first": value(pf.first_line_indent),
        "keep_next": value(pf.keep_with_next),
        "page_before": value(pf.page_break_before),
    })
    for run in p.runs:
        if not run.text:
            continue
        color = run.font.color.rgb
        fonts = run._element.rPr.rFonts if run._element.rPr is not None else None
        print("  run", {
            "text": run.text[:35],
            "name": run.font.name,
            "ascii": fonts.get(qn("w:ascii")) if fonts is not None else None,
            "size": value(run.font.size),
            "bold": run.bold,
            "italic": run.italic,
            "color": str(color) if color else None,
            "caps": run.font.all_caps,
        })

table = doc.tables[0]
print("TABLE")
print(table._tbl.tblPr.xml)
for cell in table.rows[0].cells:
    print(cell._tc.tcPr.xml)
