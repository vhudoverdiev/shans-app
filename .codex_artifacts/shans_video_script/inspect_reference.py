from pathlib import Path
from docx import Document

source = Path(r"C:\Users\Владимир\Desktop\Сайты\Передача\Презентация Peredacha\Раб.файлы для видео\Сценарий_видеопрезентации_CRM_Передача.docx")
doc = Document(source)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
for index, paragraph in enumerate(doc.paragraphs, 1):
    text = paragraph.text.replace("\n", "\\n")
    if text.strip():
        print(f"P{index:03d} [{paragraph.style.name}] {text}")

for table_index, table in enumerate(doc.tables, 1):
    print(f"TABLE {table_index}: rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows, 1):
        values = [cell.text.replace("\n", " / ") for cell in row.cells]
        print(f"  R{row_index:02d}: " + " || ".join(values))

for section_index, section in enumerate(doc.sections, 1):
    for label, part in (("HEADER", section.header), ("FOOTER", section.footer)):
        values = [p.text for p in part.paragraphs if p.text.strip()]
        print(f"{label} {section_index}: {values}")
