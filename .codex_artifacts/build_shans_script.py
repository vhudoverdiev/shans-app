from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("Shans_video_script_90sec.docx").resolve()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="DADCE0"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)


def add_run(paragraph, text, *, bold=False, size=11, color="111827"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    add_run(p, text, bold=True, size=15, color="2563EB")
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_run(p, text, size=11, color="111827")
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    add_run(title, "SHANS", bold=True, size=24, color="111827")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    add_run(subtitle, "Сценарий видеоролика на 1,5 минуты", size=13, color="64748B")

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(12)
    lead.paragraph_format.line_spacing = 1.15
    add_run(
        lead,
        "Цель ролика: быстро показать Shans как личную систему управления, где график, проекты, отчёты, обучение, спорт и уведомления собраны в одном удобном веб-приложении.",
        size=11,
        color="111827",
    )

    add_heading(doc, "План записи")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(2.7), Cm(5.5), Cm(8.0)]
    headers = ["Таймкод", "Действие на экране", "Текст диктора"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "EFF6FF")
        set_cell_border(cell, "BFDBFE")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, text, bold=True, size=10, color="2563EB")

    rows = [
        (
            "0:00-0:12",
            "Открыть главную страницу. Показать логотип, приветственный экран и основные разделы.",
            "Shans — это личная система управления, где повседневные задачи не разбросаны по заметкам, таблицам и чатам, а собраны в одном удобном веб-приложении.",
        ),
        (
            "0:12-0:30",
            "Плавно пройти по разделам: график, съёмки, фотопроекты, сценарии.",
            "В графике можно планировать дела, видеть ближайшие задачи и получать напоминания. Для съёмок есть отдельный блок: записи клиентов, фотопроекты, временные окна, оплата и подготовка сценариев.",
        ),
        (
            "0:30-0:50",
            "Показать отчёты, бюджет, автомобиль и импорт/экспорт Excel.",
            "Отчётный блок помогает контролировать финансы, баланс и историю операций. Раздел автомобиля хранит выполненные и плановые работы, а Excel-выгрузки позволяют быстро забрать данные в привычный формат.",
        ),
        (
            "0:50-1:08",
            "Показать обучение, спорт, питание, затем мобильное нижнее меню.",
            "Shans закрывает не только рабочие процессы. Внутри есть курсы IT и English, тренировки, питание и контроль прогресса. На телефоне всё адаптировано под быстрые действия в пару касаний.",
        ),
        (
            "1:08-1:22",
            "Показать уведомления, настройки аккаунта, 2FA и роли доступа.",
            "Система умеет отправлять push-уведомления, поддерживает двухфакторную защиту и роли доступа. Это делает сервис удобным не только для одного пользователя, но и для аккуратного разделения рабочих зон.",
        ),
        (
            "1:22-1:30",
            "Финальный кадр: логотип Shans и фотография разработчика.",
            "Проект создан Владимиром Худовердиевым под реальные задачи. Shans помогает держать порядок, видеть важное вовремя и освобождать внимание для работы и развития.",
        ),
    ]

    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_border(cells[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_run(p, text, bold=True, size=10, color="111827")
            else:
                add_run(p, text, size=10, color="111827" if i == 2 else "334155")

    add_heading(doc, "Ключевые особенности, которые важно подчеркнуть")
    points = [
        "Единое рабочее пространство: личные дела, проекты, отчёты и развитие в одном сервисе.",
        "Мобильный сценарий: нижнее меню, быстрые действия и push-уведомления.",
        "Практичность: бюджет, автомобиль, Excel-выгрузки и импорт-центр.",
        "Безопасность: авторизация, 2FA, роли доступа и настройки аккаунта.",
        "Авторский характер проекта: система создана под реальные задачи и может развиваться дальше.",
    ]
    for point in points:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(5)
        add_run(p, "• ", bold=True, size=11, color="2563EB")
        add_run(p, point, size=11, color="111827")

    add_heading(doc, "Финальная фраза для спокойного завершения")
    add_body(
        doc,
        "Shans — это не просто набор разделов, а личная система, которая помогает держать порядок каждый день: планировать, считать, помнить, учиться и двигаться вперёд.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
