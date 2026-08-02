import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app import create_app
from app.database import get_connection, get_master_connection, init_db
from app.models import (
    create_budget_entry,
    delete_budget_entry,
    get_all_budget_entries,
    get_balance_for_month,
    get_balance_history,
    get_budget_entry_by_id,
    get_budget_summary,
    get_current_balance,
    save_balance_history,
    set_current_balance,
    update_budget_entry,
)
from app.routes import MONTHS
from config import Config


INCOME = "Доход"
EXPENSE = "Расход"
JANUARY = MONTHS[0]
FEBRUARY = MONTHS[1]


class BudgetModelContractTests(unittest.TestCase):
    def setUp(self):
        self.original_database_name = Config.DATABASE_NAME
        self.temp_dir = tempfile.TemporaryDirectory()
        Config.DATABASE_NAME = str(Path(self.temp_dir.name) / "budget.db")
        init_db()

    def tearDown(self):
        Config.DATABASE_NAME = self.original_database_name
        self.temp_dir.cleanup()

    def test_budget_summary_counts_only_selected_month_and_entry_types(self):
        create_budget_entry(INCOME, JANUARY, "Salary", 100000)
        create_budget_entry(INCOME, JANUARY, "Gift", 5000)
        create_budget_entry(EXPENSE, JANUARY, "Food", 15000)
        create_budget_entry(EXPENSE, FEBRUARY, "Food", 8000)
        create_budget_entry("Transfer", JANUARY, "Ignored", 999)

        summary = get_budget_summary(JANUARY)

        self.assertEqual(summary["income"], 105000)
        self.assertEqual(summary["expense"], 15000)
        self.assertEqual(summary["balance"], 90000)

    def test_budget_filters_and_sorting_return_only_current_year_entries(self):
        create_budget_entry(INCOME, FEBRUARY, "Salary", 100000)
        create_budget_entry(EXPENSE, JANUARY, "Food", 5000)
        create_budget_entry(EXPENSE, JANUARY, "Auto", 15000)
        conn = get_connection()
        try:
            conn.execute(
                """
                INSERT INTO budget_entries (
                    entry_type, month_name, year_value, category, amount
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (INCOME, JANUARY, 1999, "Old", 1),
            )
            conn.commit()
        finally:
            conn.close()

        january_expenses = get_all_budget_entries(
            month_filter=JANUARY,
            type_filter=EXPENSE,
            sort_by="category",
        )
        month_sorted = get_all_budget_entries(sort_by="month")

        self.assertEqual(
            [(row["category"], row["amount"]) for row in january_expenses],
            [("Auto", 15000), ("Food", 5000)],
        )
        self.assertEqual([row["month_name"] for row in month_sorted], [JANUARY, JANUARY, FEBRUARY])
        self.assertNotIn("Old", {row["category"] for row in month_sorted})

    def test_update_and_delete_budget_entry_keep_year_current_and_remove_record(self):
        create_budget_entry(EXPENSE, JANUARY, "Food", 5000)

        update_budget_entry(1, INCOME, FEBRUARY, "Salary", 120000)
        updated = get_budget_entry_by_id(1)
        delete_budget_entry(1)

        self.assertEqual(updated["entry_type"], INCOME)
        self.assertEqual(updated["month_name"], FEBRUARY)
        self.assertEqual(updated["category"], "Salary")
        self.assertEqual(updated["amount"], 120000)
        self.assertIsNone(get_budget_entry_by_id(1))

    def test_balance_history_upserts_month_and_orders_by_calendar_month(self):
        save_balance_history(FEBRUARY, 2000, year_value=2026)
        save_balance_history(JANUARY, 1000, year_value=2026)
        save_balance_history(JANUARY, 1500, year_value=2026)

        history = get_balance_history(2026)

        self.assertEqual(
            [(row["month_name"], row["balance_value"]) for row in history],
            [(JANUARY, 1500), (FEBRUARY, 2000)],
        )
        self.assertEqual(get_balance_for_month(JANUARY, 2026), 1500)
        self.assertIsNone(get_balance_for_month(MONTHS[2], 2026))

    def test_current_balance_defaults_to_zero_and_can_be_updated(self):
        self.assertEqual(get_current_balance(), 0)

        set_current_balance(12345)

        self.assertEqual(get_current_balance(), 12345)


class BudgetRouteContractTests(unittest.TestCase):
    def setUp(self):
        self.original_database_name = Config.DATABASE_NAME
        self.original_admin_username = os.environ.get("ADMIN_USERNAME")
        self.original_admin_password = os.environ.get("ADMIN_PASSWORD")
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        Config.DATABASE_NAME = str(Path(self.temp_dir.name) / "budget-routes.db")
        os.environ["ADMIN_USERNAME"] = "admin"
        os.environ["ADMIN_PASSWORD"] = "AdminPass-2026"
        with patch.dict(os.environ, {"WERKZEUG_RUN_MAIN": "false"}):
            self.app = create_app()
        self.app.config.update(TESTING=True)

    def tearDown(self):
        Config.DATABASE_NAME = self.original_database_name
        if self.original_admin_username is None:
            os.environ.pop("ADMIN_USERNAME", None)
        else:
            os.environ["ADMIN_USERNAME"] = self.original_admin_username
        if self.original_admin_password is None:
            os.environ.pop("ADMIN_PASSWORD", None)
        else:
            os.environ["ADMIN_PASSWORD"] = self.original_admin_password
        self.temp_dir.cleanup()

    def _admin_id(self):
        conn = get_master_connection()
        try:
            row = conn.execute("SELECT id FROM users WHERE username = ?", ("admin",)).fetchone()
            return int(row["id"])
        finally:
            conn.close()

    def _login(self, client):
        with client.session_transaction() as session:
            session["_user_id"] = str(self._admin_id())
            session["_fresh"] = True
            session["_csrf_token"] = "test-token"

    def test_add_entry_route_validates_required_integer_positive_amount_and_custom_category(self):
        with self.app.test_client() as client:
            self._login(client)
            missing = client.post(
                "/budget/manage",
                data={"_csrf_token": "test-token", "form_type": "add_entry"},
            )
            non_integer = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "add_entry",
                    "month_name": JANUARY,
                    "entry_type": EXPENSE,
                    "amount": "10.5",
                    "category": "Food",
                },
            )
            negative = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "add_entry",
                    "month_name": JANUARY,
                    "entry_type": EXPENSE,
                    "amount": "-1",
                    "category": "Food",
                },
            )
            missing_custom = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "add_entry",
                    "month_name": JANUARY,
                    "entry_type": EXPENSE,
                    "amount": "100",
                    "category": "Другое",
                    "custom_category": "",
                },
            )

        self.assertEqual(missing.status_code, 302)
        self.assertEqual(non_integer.status_code, 302)
        self.assertEqual(negative.status_code, 302)
        self.assertEqual(missing_custom.status_code, 302)
        self.assertEqual(get_all_budget_entries(), [])

    def test_add_entry_route_accepts_custom_category_and_edit_route_updates_it(self):
        with self.app.test_client() as client:
            self._login(client)
            create_response = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "add_entry",
                    "month_name": JANUARY,
                    "entry_type": EXPENSE,
                    "amount": "2500",
                    "category": "Другое",
                    "custom_category": "Books",
                },
            )
            edit_response = client.post(
                "/budget/edit/1",
                data={
                    "_csrf_token": "test-token",
                    "month_name": FEBRUARY,
                    "entry_type": INCOME,
                    "amount": "5000",
                    "category": "Salary",
                },
            )

        entry = get_budget_entry_by_id(1)
        self.assertEqual(create_response.status_code, 302)
        self.assertEqual(edit_response.status_code, 302)
        self.assertEqual(entry["month_name"], FEBRUARY)
        self.assertEqual(entry["entry_type"], INCOME)
        self.assertEqual(entry["category"], "Salary")
        self.assertEqual(entry["amount"], 5000)

    def test_set_balance_route_validates_input_and_writes_current_month_history(self):
        with self.app.test_client() as client:
            self._login(client)
            invalid_response = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "set_balance",
                    "current_balance": "bad",
                },
            )
            negative_response = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "set_balance",
                    "current_balance": "-1",
                },
            )
            valid_response = client.post(
                "/budget/manage",
                data={
                    "_csrf_token": "test-token",
                    "form_type": "set_balance",
                    "current_balance": "7777",
                },
            )

        self.assertEqual(invalid_response.status_code, 302)
        self.assertEqual(negative_response.status_code, 302)
        self.assertEqual(valid_response.status_code, 302)
        self.assertEqual(get_current_balance(), 7777)
        self.assertEqual(len(get_balance_history()), 1)
        self.assertEqual(get_balance_history()[0]["balance_value"], 7777)

    def test_budget_delete_selected_ignores_malformed_ids_and_export_returns_excel(self):
        create_budget_entry(INCOME, JANUARY, "Salary", 100000)
        create_budget_entry(EXPENSE, JANUARY, "Food", 15000)

        with self.app.test_client() as client:
            self._login(client)
            delete_response = client.post(
                "/budget/delete-selected",
                data={"_csrf_token": "test-token", "selected_ids": "bad,1,999"},
            )
            export_response = client.get(f"/budget/export?month={JANUARY}")

        self.assertEqual(delete_response.status_code, 302)
        remaining = get_all_budget_entries()
        self.assertEqual(len(remaining), 1)
        self.assertEqual(remaining[0]["category"], "Food")
        self.assertEqual(export_response.status_code, 200)
        self.assertEqual(
            export_response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook_path = Path(self.temp_dir.name) / "budget-export.xlsx"
        workbook_path.write_bytes(export_response.data)
        workbook = load_workbook(workbook_path)
        self.assertGreaterEqual(len(workbook.sheetnames), 1)


if __name__ == "__main__":
    unittest.main()
