import unittest

from openpyxl import load_workbook

from app.utils import (
    build_budget_excel,
    build_photo_project_excel,
    build_scenarios_excel,
    build_shootings_excel,
)


INCOME = "Доход"
EXPENSE = "Расход"


def load_export_workbook(stream):
    return load_workbook(stream, data_only=True)


class ExcelExportContractTests(unittest.TestCase):
    def test_budget_export_preserves_summary_and_operation_rows(self):
        workbook = load_export_workbook(
            build_budget_excel(
                entries=[
                    {
                        "month_name": "January",
                        "entry_type": INCOME,
                        "category": "Salary",
                        "amount": 125000.49,
                    },
                    {
                        "month_name": "January",
                        "entry_type": EXPENSE,
                        "category": "Rent",
                        "amount": 45000.51,
                    },
                ],
                summary={"income": 125000.49, "expense": 45000.51, "balance": 79999.98},
                current_balance=180000.4,
                selected_month="January",
            )
        )

        sheet = workbook.active

        self.assertIn("January", sheet["A1"].value)
        self.assertEqual(sheet["B4"].value, 125000)
        self.assertEqual(sheet["B5"].value, 45001)
        self.assertEqual(sheet["B6"].value, 80000)
        self.assertEqual(sheet["B7"].value, 180000)
        self.assertEqual(sheet.cell(row=12, column=1).value, "January")
        self.assertEqual(sheet.cell(row=12, column=2).value, INCOME)
        self.assertEqual(sheet.cell(row=12, column=3).value, "Salary")
        self.assertEqual(sheet.cell(row=12, column=4).value, 125000)
        self.assertEqual(sheet.cell(row=13, column=2).value, EXPENSE)
        self.assertEqual(sheet.cell(row=13, column=4).value, 45001)

    def test_budget_export_keeps_headers_when_there_are_no_operations(self):
        workbook = load_export_workbook(
            build_budget_excel(
                entries=[],
                summary={"income": 0, "expense": 0, "balance": 0},
                current_balance=0,
                selected_month="February",
            )
        )

        sheet = workbook.active

        self.assertIn("February", sheet["A1"].value)
        self.assertEqual(sheet.max_row, 11)
        self.assertEqual([sheet.cell(row=11, column=col).value for col in range(1, 5)], [
            "Месяц",
            "Тип",
            "Категория",
            "Сумма",
        ])

    def test_shootings_export_writes_payment_contract_and_caps_remaining_at_zero(self):
        workbook = load_export_workbook(
            build_shootings_excel(
                shootings=[
                    {
                        "shooting_date_display": "10.08.2099",
                        "shooting_time": "14:30",
                        "project_name": "Wedding",
                        "client_name": "Anna",
                        "phone": "+70000000000",
                        "duration_hours": 3,
                        "price": "50000.6",
                        "prepayment": "15000.2",
                        "notes": "Outdoor",
                    },
                    {
                        "shooting_date": "2099-08-11",
                        "project_name": "Studio",
                        "price": 10000,
                        "prepayment": 25000,
                    },
                ],
                report_title="Shootings report",
            )
        )

        sheet = workbook.active

        self.assertEqual(sheet["A1"].value, "Shootings report")
        self.assertIn("2", sheet["A3"].value)
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertEqual(sheet.cell(row=6, column=1).value, "10.08.2099")
        self.assertEqual(sheet.cell(row=6, column=7).value, 50001)
        self.assertEqual(sheet.cell(row=6, column=8).value, 15000)
        self.assertEqual(sheet.cell(row=6, column=9).value, 35001)
        self.assertEqual(sheet.cell(row=7, column=1).value, "2099-08-11")
        self.assertEqual(sheet.cell(row=7, column=9).value, 0)
        self.assertTrue(sheet.cell(row=7, column=4).value)

    def test_shootings_export_treats_missing_money_as_zero(self):
        workbook = load_export_workbook(
            build_shootings_excel(
                shootings=[
                    {
                        "project_name": "No payment yet",
                        "client_name": "Client",
                    }
                ],
                report_title="No payments",
            )
        )

        sheet = workbook.active

        self.assertEqual(sheet.cell(row=6, column=7).value, 0)
        self.assertEqual(sheet.cell(row=6, column=8).value, 0)
        self.assertEqual(sheet.cell(row=6, column=9).value, 0)
        self.assertTrue(sheet.cell(row=6, column=1).value)
        self.assertTrue(sheet.cell(row=6, column=10).value)

    def test_photo_project_export_includes_project_metadata_and_booking_finances(self):
        workbook = load_export_workbook(
            build_photo_project_excel(
                project={
                    "title": "Spring mini sessions",
                    "city": "Moscow",
                    "project_date_display": "01.05.2099",
                    "time_range_display": "10:00-18:00",
                    "address": "Studio 1",
                },
                bookings=[
                    {
                        "client_name": "Olga",
                        "client_contact": "@olga",
                        "booking_date_display": "01.05.2099",
                        "booking_time": "10:30",
                        "duration_minutes": 45,
                        "makeup_start_time": "09:45",
                        "price": "12000",
                        "prepayment": "4000",
                        "status": "confirmed",
                        "comment": "Bring red dress",
                    }
                ],
            )
        )

        sheet = workbook.active

        self.assertIn("Spring mini sessions", sheet["A1"].value)
        self.assertIn("Moscow", sheet["A3"].value)
        self.assertIn("01.05.2099", sheet["A3"].value)
        self.assertIn("Studio 1", sheet["A4"].value)
        self.assertIn("1", sheet["A5"].value)
        self.assertEqual(sheet.freeze_panes, "A8")
        self.assertEqual(sheet.cell(row=8, column=1).value, "Olga")
        self.assertEqual(sheet.cell(row=8, column=5).value, "45 мин")
        self.assertEqual(sheet.cell(row=8, column=7).value, 12000)
        self.assertEqual(sheet.cell(row=8, column=8).value, 4000)
        self.assertEqual(sheet.cell(row=8, column=9).value, 8000)
        self.assertEqual(sheet.cell(row=8, column=10).value, "confirmed")

    def test_photo_project_export_uses_safe_defaults_for_sparse_booking(self):
        workbook = load_export_workbook(
            build_photo_project_excel(
                project={},
                bookings=[
                    {
                        "price": None,
                        "prepayment": None,
                    }
                ],
            )
        )

        sheet = workbook.active

        self.assertTrue(sheet["A1"].value)
        self.assertTrue(sheet["A3"].value)
        self.assertEqual(sheet.cell(row=8, column=5).value, "15 мин")
        self.assertEqual(sheet.cell(row=8, column=7).value, 0)
        self.assertEqual(sheet.cell(row=8, column=8).value, 0)
        self.assertEqual(sheet.cell(row=8, column=9).value, 0)
        self.assertTrue(sheet.cell(row=8, column=1).value)

    def test_scenarios_export_preserves_status_and_long_text_contract(self):
        workbook = load_export_workbook(
            build_scenarios_excel(
                scenarios=[
                    {
                        "title": "Morning reel",
                        "shooting_date_display": "15.09.2099",
                        "effective_status_label": "Ready",
                        "scenario_text": "Shot 1\nShot 2\nFinal CTA",
                    },
                    {
                        "title": "Draft reel",
                        "shooting_date": "2099-09-16",
                        "scenario_text": "",
                    },
                ],
                report_title="Scenario report",
            )
        )

        sheet = workbook.active

        self.assertEqual(sheet["A1"].value, "Scenario report")
        self.assertEqual(sheet.freeze_panes, "A5")
        self.assertEqual(sheet.cell(row=5, column=1).value, "Morning reel")
        self.assertEqual(sheet.cell(row=5, column=2).value, "15.09.2099")
        self.assertEqual(sheet.cell(row=5, column=3).value, "Ready")
        self.assertIn("Final CTA", sheet.cell(row=5, column=4).value)
        self.assertEqual(sheet.cell(row=6, column=2).value, "2099-09-16")
        self.assertTrue(sheet.cell(row=6, column=3).value)
        self.assertTrue(sheet.cell(row=6, column=4).value)

    def test_scenarios_export_keeps_header_only_report_valid(self):
        workbook = load_export_workbook(build_scenarios_excel([], "Empty scenarios"))

        sheet = workbook.active

        self.assertEqual(sheet["A1"].value, "Empty scenarios")
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.cell(row=4, column=1).value, "Название")
        self.assertEqual(sheet.cell(row=4, column=4).value, "Сценарий")


if __name__ == "__main__":
    unittest.main()
